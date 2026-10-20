"""AI 報告的版面：內容先組成 blocks，再分別輸出成 .docx 與 .pdf，兩份內容一定一致。

PDF 用 fpdf2 直接排版（不經 LibreOffice 轉檔，映像檔不用多裝幾百 MB），中文字型是
Noto Sans TC（SIL OFL），由 Dockerfile 下載到 REPORT_FONT_DIR 並嵌入 PDF，任何裝置開啟都看得到中文。"""

import io
import logging
import os
from dataclasses import dataclass

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from fpdf import FPDF
from fpdf.enums import WrapMode
from fpdf.fonts import FontFace

# fpdf2 嵌入字型時 fontTools 每次裁切都印一堆 INFO，main.py 把 root 設成 INFO 會全進 Cloud Logging
logging.getLogger("fontTools").setLevel(logging.WARNING)

FONT_DIR = os.environ.get("REPORT_FONT_DIR", "/usr/share/fonts/noto-tc")
FONT_REGULAR = os.path.join(FONT_DIR, "NotoSansTC-Regular.otf")
FONT_BOLD = os.path.join(FONT_DIR, "NotoSansTC-Bold.otf")

# 圖表 PNG 的解析度（charts.py 存檔用同一個值）。圖片照原尺寸放、不拉滿版寬，
# 每張圖的字級才會一致（標籤短的圖裁切後比較窄，拉滿就會被放大）。
IMAGE_DPI = 200

_INK = (11, 11, 11)
_INK_SOFT = (82, 81, 78)
_HEADER_FILL = (234, 241, 250)


@dataclass
class Heading:
    text: str
    level: int = 2


@dataclass
class Paragraph:
    text: str
    small: bool = False  # 圖說、AI 免責聲明這類次要文字


@dataclass
class Bullet:
    text: str


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    col_widths: tuple[int, ...] | None = None  # 欄寬比例，None＝平均


@dataclass
class Image:
    png: bytes
    caption: str

    @property
    def width_mm(self) -> float:
        width_px = int.from_bytes(self.png[16:20], "big")  # PNG IHDR 的寬度欄位
        return width_px / IMAGE_DPI * 25.4


Block = Heading | Paragraph | Bullet | Table | Image


def render_docx(blocks: list[Block]) -> bytes:
    doc = Document()
    for b in blocks:
        if isinstance(b, Heading):
            doc.add_heading(b.text, level=b.level)
        elif isinstance(b, Paragraph):
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            if b.small:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(*_INK_SOFT)
        elif isinstance(b, Bullet):
            doc.add_paragraph(b.text, style="List Bullet")
        elif isinstance(b, Table):
            table = doc.add_table(rows=1, cols=len(b.headers))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, b.headers):
                cell.text = text
            for row in b.rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
        elif isinstance(b, Image):
            doc.add_picture(io.BytesIO(b.png), width=Cm(min(15, b.width_mm / 10)))
            run = doc.add_paragraph().add_run(b.caption)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(*_INK_SOFT)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _PDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("NotoTC", size=8)
        self.set_text_color(*_INK_SOFT)
        self.cell(0, 5, str(self.page_no()), align="C")


_HEADING_SIZES = {1: 18, 2: 14, 3: 12}


def render_pdf(blocks: list[Block]) -> bytes:
    pdf = _PDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("NotoTC", "", FONT_REGULAR)
    pdf.add_font("NotoTC", "B", FONT_BOLD)
    pdf.add_page()
    pdf.set_text_color(*_INK)

    def text(value: str, size: float, style: str = "", indent: float = 0, color=_INK):
        pdf.set_font("NotoTC", style, size)
        pdf.set_text_color(*color)
        pdf.set_x(pdf.l_margin + indent)
        # 中文沒有空白可斷行，要用 CHAR 模式才會在任意字之間換行
        pdf.multi_cell(pdf.epw - indent, size * 0.55, value, wrapmode=WrapMode.CHAR, new_x="LMARGIN", new_y="NEXT")

    for i, b in enumerate(blocks):
        if isinstance(b, Heading):
            # 標題下方放不下約 4 行內容就直接換頁，避免標題孤零零留在頁尾
            if i and pdf.get_y() > pdf.page_break_trigger - 35:
                pdf.add_page()
            elif i:
                pdf.ln(4)
            text(b.text, _HEADING_SIZES.get(b.level, 12), "B")
            pdf.ln(1.5)
        elif isinstance(b, Paragraph):
            text(b.text, 9 if b.small else 11, color=_INK_SOFT if b.small else _INK)
            pdf.ln(1.5)
        elif isinstance(b, Bullet):
            pdf.set_font("NotoTC", "", 11)
            y = pdf.get_y()
            pdf.set_xy(pdf.l_margin + 1, y)
            pdf.cell(4, 11 * 0.55, "•")
            pdf.set_xy(pdf.l_margin, y)
            text(b.text, 11, indent=5)
            pdf.ln(0.8)
        elif isinstance(b, Table):
            pdf.set_font("NotoTC", "", 10)
            pdf.set_text_color(*_INK)
            with pdf.table(
                col_widths=b.col_widths,
                headings_style=FontFace(emphasis="BOLD", fill_color=_HEADER_FILL),
                line_height=6,
                text_align="LEFT",
                wrapmode=WrapMode.CHAR,
            ) as table:
                for row in [b.headers, *b.rows]:
                    r = table.row()
                    for cell in row:
                        r.cell(cell)
            pdf.ln(2)
        elif isinstance(b, Image):
            pdf.ln(1)
            pdf.image(io.BytesIO(b.png), w=min(pdf.epw, b.width_mm))
            text(b.caption, 9, color=_INK_SOFT)
            pdf.ln(2)
    return bytes(pdf.output())
