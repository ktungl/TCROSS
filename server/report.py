"""Phase 3b: call Vertex AI Gemini on a GenerationJob's uploaded source files,
then assemble the structured result into a real .docx report. See ROADMAP.md
「GCP 整合進度」Phase 3b for the architecture this implements."""

import io
import json
import logging
import mimetypes
import os

from docx import Document
from google import genai
from google.genai import types

import charts
from document import Block, Bullet, Heading, Image, Paragraph, Table
from storage import download_bytes

GCP_PROJECT = os.environ["GCP_PROJECT"]
GCP_LOCATION = os.environ.get("GCP_LOCATION", "us-central1")
GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")
GCS_BUCKET = os.environ["GCS_BUCKET"]
# 照片／影片送進 Gemini 的解析度：LOW 每張圖／每格影格的 token 數只有預設的約 1/4，
# 大幅壓低費用；如果發現簽到表之類小字辨識不出來，再把環境變數改成 MEDIUM／HIGH。
GEMINI_MEDIA_RESOLUTION = os.environ.get("GEMINI_MEDIA_RESOLUTION", "LOW").upper()

_client: genai.Client | None = None


def _get_client() -> genai.Client:
    """Cached per Cloud Run instance, same pattern as storage._client() — picks
    up the runtime service account's ADC automatically (needs roles/aiplatform.user
    granted on it, see server/README.md)."""
    global _client
    if _client is None:
        _client = genai.Client(vertexai=True, project=GCP_PROJECT, location=GCP_LOCATION)
    return _client


_MIME_FALLBACK_BY_FOLDER = {
    "photo": "image/jpeg",
    "audio": "audio/mpeg",
    "video": "video/mp4",
    "doc": "application/pdf",
}


def _guess_mime_type(object_path: str) -> str:
    guessed, _ = mimetypes.guess_type(object_path)
    if guessed:
        return guessed
    # Fallback by the ALLOWED_FOLDERS segment (activities/{id}/{folder}/{filename})
    # for extensionless or unrecognized filenames.
    parts = object_path.split("/")
    folder = parts[2] if len(parts) > 2 else ""
    return _MIME_FALLBACK_BY_FOLDER.get(folder, "application/octet-stream")


_RESPONSE_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "summary": {"type": "STRING"},
        "highlights": {"type": "ARRAY", "items": {"type": "STRING"}},
        "kpis": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "k": {"type": "STRING"},
                    "v": {"type": "STRING"},
                    "u": {"type": "STRING"},
                },
                "required": ["k", "v"],
            },
        },
    },
    "required": ["summary", "highlights", "kpis"],
}

_PROMPT = (
    "你是社福機構的活動成果報告助手。以下是一場活動上傳的原始素材（照片／錄音／影片／文件），"
    "請閱讀並歸納，用繁體中文回覆：\n"
    "1. summary：活動內容簡述與效益（2-4 句話）\n"
    "2. highlights：3-6 條重點條列\n"
    "3. kpis：能從素材中直接辨識出的量化成果指標（例如參加人數、場次數），"
    "k=項目名稱、v=數值、u=單位；素材中看不出可量化的指標就回傳空陣列，不要杜撰數字。"
)


def generate_json(parts: list, schema: dict, **config) -> dict:
    """Call Gemini with structured-JSON output; shared by report / assist / rewrite."""
    response = _get_client().models.generate_content(
        model=GEMINI_MODEL,
        contents=types.Content(role="user", parts=parts),
        config=types.GenerateContentConfig(
            response_mime_type="application/json", response_schema=schema, **config
        ),
    )
    # 被安全過濾擋下或輸出被截斷時 response.text 是 None，直接 json.loads 只會留下看不懂的 TypeError
    if not response.text:
        raise ValueError("AI 沒有回傳內容（可能被內容過濾擋下），請換一份素材再試")
    return json.loads(response.text)


def _docx_text(raw: bytes) -> str:
    """Word 檔的文字：段落在前、表格逐列接在後面（Gemini 只需要內容，不需要版面）。"""
    doc = Document(io.BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _source_part(path: str) -> types.Part:
    """Gemini 不收 .docx（會回 400「The document has no pages」），先在這裡取出文字；
    其他格式直接用 GCS URI 給 Gemini 讀，不必先下載到這個服務。"""
    # 看副檔名，不靠 mimetypes：python:3.12-slim 沒有 /etc/mime.types，認不出 .docx，
    # 會落到資料夾預設的 application/pdf
    if path.lower().endswith(".docx"):
        name = path.rsplit("/", 1)[-1]
        text = _docx_text(download_bytes(path)) or "（這份 Word 檔沒有文字內容）"
        return types.Part.from_text(text=f"【文件：{name}】\n{text}")
    return types.Part.from_uri(file_uri=f"gs://{GCS_BUCKET}/{path}", mime_type=_guess_mime_type(path))


def _generate_json(object_paths: list[str], prompt: str, schema: dict, **config) -> dict:
    """Feed the job's uploaded files to Gemini (see _source_part) and ask for
    structured JSON back."""
    parts = [_source_part(path) for path in object_paths]
    parts.append(types.Part.from_text(text=prompt))
    return generate_json(parts, schema, **config)


def analyze_sources(object_paths: list[str]) -> dict:
    if not object_paths:
        raise ValueError("沒有任何素材檔案")
    data = _generate_json(
        object_paths, _PROMPT, _RESPONSE_SCHEMA,
        media_resolution=f"MEDIA_RESOLUTION_{GEMINI_MEDIA_RESOLUTION}",
    )
    return {
        "summary": str(data.get("summary") or ""),
        "highlights": [str(h) for h in (data.get("highlights") or [])],
        "kpis": [
            {"k": str(k.get("k", "")), "v": str(k.get("v", "")), "u": str(k.get("u", ""))}
            for k in (data.get("kpis") or [])
        ],
    }


def _safe_charts(build) -> list[tuple[bytes, str]]:
    """圖表只是輔助：畫不出來（字型、資料格式）就略過，不讓整份報告失敗。"""
    try:
        result = build()
    except Exception:  # noqa: BLE001
        logging.getLogger("tcross.audit").exception("report chart failed")
        return []
    if result is None:
        return []
    return result if isinstance(result, list) else [result]


def build_report(activity: dict, analysis: dict, plan_activities: list[tuple[str, list[dict]]]) -> list[Block]:
    """成果報告。plan_activities 是 [(計畫名稱, 該計畫的活動清單)]，用來畫同計畫比較圖。"""
    blocks: list[Block] = [Heading(activity.get("name") or "活動成果報告", 1)]

    blocks += [Heading("摘要"), Paragraph(analysis["summary"] or "（無摘要）")]

    if analysis["highlights"]:
        blocks.append(Heading("重點"))
        blocks += [Bullet(item) for item in analysis["highlights"]]

    if analysis["kpis"]:
        blocks.append(Heading("成果指標"))
        blocks.append(Table(["項目", "數值", "單位"], [[k["k"], k["v"], k["u"]] for k in analysis["kpis"]], (3, 1, 1)))
        blocks += [Image(png, caption) for png, caption in _safe_charts(lambda: charts.kpi_charts(analysis["kpis"]))]

    participation = _safe_charts(lambda: charts.gender_chart(activity))
    for plan_name, acts in plan_activities:
        participation += _safe_charts(lambda: charts.plan_comparison_chart(plan_name, activity.get("objectId", ""), acts))
    if participation:
        blocks.append(Heading("參與情形"))
        blocks += [Image(png, caption) for png, caption in participation]

    return blocks


# ── 會議紀錄（AI_UPGRADE_PLAN.md 1-3）──────────────────────────

_MEETING_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "title": {"type": "STRING"},
        "participants": {"type": "ARRAY", "items": {"type": "STRING"}},
        "summary": {"type": "STRING"},
        "discussions": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "topic": {"type": "STRING"},
                    "points": {"type": "ARRAY", "items": {"type": "STRING"}},
                },
                "required": ["topic", "points"],
            },
        },
        "decisions": {"type": "ARRAY", "items": {"type": "STRING"}},
        "todos": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "task": {"type": "STRING"},
                    "owner": {"type": "STRING"},
                    "due": {"type": "STRING"},
                },
                "required": ["task", "owner", "due"],
            },
        },
        "notes": {"type": "STRING"},
    },
    "required": ["title", "participants", "summary", "discussions", "decisions", "todos", "notes"],
}

_MEETING_PROMPT = (
    "以上是一場會議的錄音（多個檔案時依序接續）。請整理成會議紀錄，用繁體中文回覆，"
    "只能根據錄音內容，聽不清楚或沒提到的不要杜撰：\n"
    "- title：會議名稱或主題（錄音中沒提到就依內容寫一句話主題）\n"
    "- participants：錄音中有自我介紹或被點名的與會者（姓名或職稱），聽不出來就回傳空陣列，不要猜\n"
    "- summary：3-5 句說明這場會議談了什麼、得到什麼結論\n"
    "- discussions：依議題整理討論重點，每個議題 2-5 點\n"
    "- decisions：明確做成的決議，沒有就回傳空陣列\n"
    "- todos：會後待辦；task 寫要做什麼，owner 寫負責人（錄音中有指定才填，否則空字串），"
    "due 寫期限（照錄音中的說法，例如「下週三前」「10 月底」，沒提到就空字串）\n"
    "- notes：錄音品質不佳、中斷或有聽不清楚的段落時簡短說明，沒有就空字串"
)


def analyze_meeting(object_paths: list[str]) -> dict:
    """錄音直接以 GCS URI 給 Gemini（同 analyze_sources），只回結構化紀錄、不做逐字稿，輸出 token 少。"""
    if not object_paths:
        raise ValueError("沒有任何錄音檔")
    data = _generate_json(object_paths, _MEETING_PROMPT, _MEETING_SCHEMA)

    def strings(values) -> list[str]:
        return [s for s in (str(v).strip() for v in values or []) if s]

    return {
        "title": str(data.get("title") or "").strip(),
        "participants": strings(data.get("participants")),
        "summary": str(data.get("summary") or "").strip(),
        "discussions": [
            {"topic": str(d.get("topic") or "").strip(), "points": strings(d.get("points"))}
            for d in data.get("discussions") or []
            if str(d.get("topic") or "").strip() or d.get("points")
        ],
        "decisions": strings(data.get("decisions")),
        "todos": [
            {
                "task": str(t.get("task") or "").strip(),
                "owner": str(t.get("owner") or "").strip(),
                "due": str(t.get("due") or "").strip(),
            }
            for t in data.get("todos") or []
            if str(t.get("task") or "").strip()
        ],
        "notes": str(data.get("notes") or "").strip(),
    }


def build_meeting(activity_name: str, activity_date: str, minutes: dict) -> list[Block]:
    blocks: list[Block] = [Heading(minutes["title"] or activity_name or "會議紀錄", 1)]
    meta = "　".join(s for s in [f"活動：{activity_name}" if activity_name else "", f"日期：{activity_date}" if activity_date else ""] if s)
    if meta:
        blocks.append(Paragraph(meta))
    if minutes["participants"]:
        blocks.append(Paragraph(f"與會者（錄音中提到）：{'、'.join(minutes['participants'])}"))

    blocks += [Heading("會議摘要"), Paragraph(minutes["summary"] or "（無摘要）")]

    if minutes["discussions"]:
        blocks.append(Heading("討論重點"))
        for d in minutes["discussions"]:
            if d["topic"]:
                blocks.append(Heading(d["topic"], 3))
            blocks += [Bullet(point) for point in d["points"]]

    blocks.append(Heading("決議事項"))
    if minutes["decisions"]:
        blocks += [Paragraph(f"{i}. {item}") for i, item in enumerate(minutes["decisions"], 1)]
    else:
        blocks.append(Paragraph("（錄音中沒有明確決議）"))

    blocks.append(Heading("待辦事項"))
    if minutes["todos"]:
        blocks.append(Table(
            ["事項", "負責人", "期限"],
            [[t["task"], t["owner"] or "（未指定）", t["due"] or "—"] for t in minutes["todos"]],
            (3, 1, 1),
        ))
    else:
        blocks.append(Paragraph("（錄音中沒有明確待辦）"))

    if minutes["notes"]:
        blocks.append(Paragraph(f"錄音狀況：{minutes['notes']}"))
    blocks.append(Paragraph("※ 本紀錄由 AI 依錄音整理，決議、負責人與期限請與會議主持人確認後再發出。", small=True))
    return blocks
